from datetime import datetime
import uuid

from flask import Blueprint, render_template, redirect, url_for, request, flash, jsonify, abort, Response, current_app
from flask_login import login_required, current_user
from openpyxl import load_workbook

from models import VerificationRequest, db, User, City, Workspace, Client, WorkSite, WorkerAccess, Task
from sqlalchemy import or_

from utils.invite_utils import DEFAULT_WORKSPACE_ID, generate_invite_key
from utils.workspace_rbac import (
    can_create_workspace,
    can_extend_workspace_subscription,
    can_terminate_workspace_admin_sessions,
    can_view_workspace,
    is_platform_creator,
    workspace_admin_context_for_template,
    workspace_panel_permissions,
    workspaces_list_query,
)
from utils.crypto import decrypt_data, encrypt_data
from utils.worker_stats import get_workers_stats_batch
from utils.media_urls import avatar_url_for
from utils.workspace_admin import (
    admin_terminate_session,
    display_workspace_name,
    fix_corrupted_workspace_names,
    normalize_workspace_name,
    serialize_workspace_detail,
    workspace_card_stats,
)
from utils.workspace_utils import (
    ADMIN_STAFF_ROLES,
    CLIENT_ADMIN_ROLES,
    USER_MANAGEMENT_ROLES,
    count_admin_staff,
    get_workspace_id,
    require_workspace,
    would_exceed_admin_limit,
)

admin_bp = Blueprint('admin', __name__)

MESSENGER_STAFF_ROLES = ['creator', 'director', 'senior_dispatcher', 'dispatcher', 'manager']
WORKER_ROLES = ['worker', 'brigadir']

NAME_COLUMNS = {'имя/организация', 'имя', 'организация', 'name', 'фio', 'фio/организация', 'фio / организация'}
PHONE_COLUMNS = {'телефон', 'phone', 'tel', 'тел'}
ADDRESS_COLUMNS = {'адрес', 'address'}
CLIENT_STATUS_ACTIVE = 'active'
CLIENT_STATUS_FROZEN = 'frozen'


def _normalize_allowed_locations(raw):
    if raw is None:
        return []
    if isinstance(raw, str):
        import json
        try:
            raw = json.loads(raw) if raw else []
        except (TypeError, ValueError):
            return []
    if not isinstance(raw, list):
        return []
    return [int(x) for x in raw if x is not None]


def _append_city_to_allowed_locations(user, city_id):
    if not city_id:
        return
    locations = _normalize_allowed_locations(user.allowed_locations)
    city_id = int(city_id)
    if city_id not in locations:
        locations.append(city_id)
    user.allowed_locations = locations


VERIFICATION_MODERATOR_ROLES = ('creator', 'director')
CITY_MODERATOR_ROLES = ('senior_dispatcher', 'dispatcher')
MODERATION_BADGE_ROLES = VERIFICATION_MODERATOR_ROLES + CITY_MODERATOR_ROLES
PENDING_VERIFICATION_STATUS = 'pending'


def moderation_city_room(city_id):
    return f'moderation_city_{int(city_id)}'


def _user_moderation_city_ids(user):
    return _normalize_allowed_locations(getattr(user, 'allowed_locations', None))


def join_user_moderation_rooms(user, join_room_fn):
    """Подключает диспетчера к socket-комнатам модерации по его городам."""
    if not user or not getattr(user, 'is_authenticated', False):
        return
    if user.role not in CITY_MODERATOR_ROLES:
        return
    for city_id in _user_moderation_city_ids(user):
        join_room_fn(moderation_city_room(city_id))


def _can_moderate_verifications():
    if current_user.platform_role == 'super_admin':
        return True
    if current_user.role in VERIFICATION_MODERATOR_ROLES:
        return True
    if current_user.role in CITY_MODERATOR_ROLES:
        return bool(_user_moderation_city_ids(current_user))
    return False


def _can_view_moderation_badge():
    if current_user.platform_role == 'super_admin':
        return True
    return current_user.role in MODERATION_BADGE_ROLES


def _scoped_verification_query(status=PENDING_VERIFICATION_STATUS):
    query = VerificationRequest.query.filter_by(status=status)
    if current_user.platform_role == 'super_admin':
        return query
    if current_user.role == 'creator':
        return query.filter(
            or_(
                VerificationRequest.workspace_id == DEFAULT_WORKSPACE_ID,
                VerificationRequest.workspace_id.is_(None),
            )
        )
    if current_user.role == 'director':
        ws_id = get_workspace_id(current_user)
        if not ws_id:
            return query.filter(False)
        return query.join(User, VerificationRequest.user_id == User.id).filter(
            User.workspace_id == ws_id,
            VerificationRequest.workspace_id == ws_id,
        )
    if current_user.role in CITY_MODERATOR_ROLES:
        city_ids = _user_moderation_city_ids(current_user)
        if not city_ids:
            return query.filter(False)
        ws_id = get_workspace_id(current_user)
        query = query.filter(VerificationRequest.city_id.in_(city_ids))
        if ws_id:
            query = query.filter(VerificationRequest.workspace_id == ws_id)
        return query
    return query.filter(False)


def _can_moderate_verification(req):
    if current_user.platform_role == 'super_admin':
        return True
    if current_user.role == 'creator':
        ws = req.workspace_id
        return ws is None or ws == DEFAULT_WORKSPACE_ID
    if current_user.role == 'director':
        ws_id = get_workspace_id(current_user)
        return bool(ws_id and req.workspace_id == ws_id)
    if current_user.role in CITY_MODERATOR_ROLES:
        city_ids = _user_moderation_city_ids(current_user)
        if not req.city_id or int(req.city_id) not in city_ids:
            return False
        ws_id = get_workspace_id(current_user)
        return not ws_id or req.workspace_id == ws_id
    return False


def _latest_verification_with_documents(user_id):
    return (
        VerificationRequest.query.filter_by(user_id=user_id)
        .order_by(VerificationRequest.id.desc())
        .first()
    )


def _can_view_verification_documents(req):
    """Просмотр зашифрованных документов анкеты (модерация + карточка исполнителя)."""
    if _can_moderate_verification(req):
        return True
    from utils.user_profiles import can_view_worker_card

    user = User.query.get(req.user_id)
    if not user or not can_view_worker_card(current_user, user):
        return False
    ws_id = get_workspace_id(current_user)
    if ws_id:
        return req.workspace_id == ws_id
    if current_user.role == 'creator':
        return req.workspace_id is None or req.workspace_id == DEFAULT_WORKSPACE_ID
    return False


def role_required(allowed_roles):
    def decorator(func):
        def wrapper(*args, **kwargs):
            if not current_user.is_authenticated:
                flash('Требуется авторизация')
                return redirect(url_for('auth.login'))
            if current_user.role not in allowed_roles:
                flash('У вас нет прав для доступа к этой странице')
                return redirect(url_for('tasks.tasks_list'))
            return func(*args, **kwargs)
        wrapper.__name__ = func.__name__
        return wrapper
    return decorator


def _parse_expires_at(raw):
    if not raw or not str(raw).strip():
        return None
    value = str(raw).strip()
    for fmt in ('%Y-%m-%d', '%d.%m.%Y', '%Y-%m-%dT%H:%M', '%d.%m.%Y %H:%M'):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def _normalize_phone(raw):
    if raw is None:
        return None
    phone = str(raw).strip()
    if not phone:
        return None
    digits = ''.join(ch for ch in phone if ch.isdigit() or ch == '+')
    if len(digits.replace('+', '')) < 10:
        return None
    return phone


def _get_workspace_client(client_id, ws_id):
    client = Client.query.get(client_id)
    if not client or client.workspace_id != ws_id:
        return None, ('Клиент не найден', 404)
    return client, None


def _delete_client_record(client):
    Task.query.filter_by(client_id=client.id).update(
        {Task.client_id: None, Task.worksite_id: None},
        synchronize_session=False,
    )
    WorkerAccess.query.filter_by(client_id=client.id).delete(synchronize_session=False)
    WorkSite.query.filter_by(client_id=client.id).delete(synchronize_session=False)
    db.session.delete(client)


def _client_phone_display(client):
    if not client.phone:
        return ''
    return decrypt_data(client.phone) or ''


def _client_json(client):
    return {
        'id': client.id,
        'name': client.name,
        'phone': _client_phone_display(client),
        'address': client.address or '',
        'status': client.status or CLIENT_STATUS_ACTIVE,
    }


def _normalize_header(value):
    return str(value or '').strip().lower().replace('ё', 'е')


def _column_map(header_row):
    mapping = {}
    for idx, cell in enumerate(header_row):
        key = _normalize_header(cell)
        if key in NAME_COLUMNS:
            mapping['name'] = idx
        elif key in PHONE_COLUMNS:
            mapping['phone'] = idx
        elif key in ADDRESS_COLUMNS:
            mapping['address'] = idx
    return mapping


def _parse_clients_workbook(file_storage):
    wb = load_workbook(file_storage, read_only=True, data_only=True)
    sheet = wb.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], 'Файл пуст'

    col_map = _column_map(rows[0])
    if 'name' not in col_map:
        return [], 'Не найдена колонка «Имя/Организация»'

    clients = []
    for row in rows[1:]:
        if not row:
            continue
        name = str(row[col_map['name']] or '').strip() if col_map['name'] < len(row) else ''
        if not name:
            continue
        phone = None
        address = None
        if 'phone' in col_map and col_map['phone'] < len(row):
            phone = _normalize_phone(row[col_map['phone']])
        if 'address' in col_map and col_map['address'] < len(row):
            address = str(row[col_map['address']] or '').strip() or None
        clients.append({'name': name, 'phone': phone, 'address': address})
    return clients, None


def _admin_context(active_tab, **extra):
    ctx = {
        'active_tab': 'admin',
        'active_admin_tab': active_tab,
        'is_creator': current_user.role == 'creator',
    }
    ctx.update(extra)
    return ctx


def _admin_redirect_for_role(role):
    if role in WORKER_ROLES:
        return redirect(url_for('admin.manage_workers'))
    return redirect(url_for('admin.manage_users'))


@admin_bp.route('/admin')
@login_required
def admin_index():
    return redirect(url_for('admin.manage_users'))


@admin_bp.route('/admin/users')
@login_required
@role_required(list(USER_MANAGEMENT_ROLES))
def manage_users():
    ws_id = get_workspace_id(current_user)
    users_query = User.query.filter(User.role.in_(MESSENGER_STAFF_ROLES))
    if ws_id:
        users_query = users_query.filter(User.workspace_id == ws_id)
    users = users_query.order_by(User.id.desc()).all()

    return render_template(
        'admin/users.html',
        **_admin_context('users', users=users),
    )


@admin_bp.route('/admin/workers')
@login_required
@role_required(list(USER_MANAGEMENT_ROLES))
def manage_workers():
    ws_id = get_workspace_id(current_user)
    workers_query = User.query.filter(User.role.in_(WORKER_ROLES))
    if ws_id:
        workers_query = workers_query.filter(User.workspace_id == ws_id)
    workers = workers_query.order_by(User.id.desc()).all()
    all_cities = City.query.order_by(City.name).all()

    worker_stats_map = get_workers_stats_batch(ws_id, [w.id for w in workers])

    pending_q = VerificationRequest.query.filter_by(status='pending')
    if ws_id:
        pending_q = pending_q.filter(
            or_(VerificationRequest.workspace_id == ws_id, VerificationRequest.workspace_id.is_(None))
        )
    pending_by_user = {r.user_id: r for r in pending_q.all()}

    return render_template(
        'admin/workers.html',
        **_admin_context(
            'workers',
            workers=workers,
            cities=all_cities,
            pending_by_user=pending_by_user,
            worker_stats_map=worker_stats_map,
        ),
    )


@admin_bp.route('/admin/workspaces')
@login_required
@role_required(['creator', 'director'])
def manage_workspaces():
    rbac_ctx = workspace_admin_context_for_template(current_user)

    if current_user.role == 'director':
        ws_id, err = require_workspace(current_user)
        if err:
            flash(err)
            return redirect(url_for('tasks.tasks_list'))

    if is_platform_creator(current_user):
        fix_corrupted_workspace_names(commit=True)

    workspaces = workspaces_list_query(current_user).all()
    workspace_stats = {}
    for ws in workspaces:
        workspace_stats[ws.id] = workspace_card_stats(ws.id, ws.admin_limit)

    share_join_url = None
    if rbac_ctx.get('show_employee_invite_panel') and current_user.workspace_id:
        director_ws = Workspace.query.get(current_user.workspace_id)
        if director_ws and director_ws.share_token:
            share_join_url = url_for(
                'auth.join_via_share_token',
                token=director_ws.share_token,
                _external=True,
            )

    return render_template(
        'admin/workspaces.html',
        **_admin_context(
            'workspaces',
            workspaces=workspaces,
            workspace_stats=workspace_stats,
            display_workspace_name=display_workspace_name,
            share_join_url=share_join_url,
            **rbac_ctx,
        ),
    )


def _workspace_access_denied_json():
    return jsonify({'status': 'error', 'error': 'Нет доступа к этому пространству'}), 403


def _creator_only_json():
    if is_platform_creator(current_user):
        return None
    return jsonify({'status': 'error', 'error': 'Доступ только для создателя платформы'}), 403


def _admin_avatar_url(user):
    return avatar_url_for(user, lambda p: url_for('uploaded_file', filename=p))


@admin_bp.route('/api/admin/workspaces/fix-names', methods=['POST'])
@login_required
@role_required(['creator'])
def api_workspace_fix_names():
    denied = _creator_only_json()
    if denied:
        return denied
    changed = fix_corrupted_workspace_names(commit=True)
    return jsonify({'status': 'success', 'fixed': changed}), 200


@admin_bp.route('/api/admin/workspaces/<int:ws_id>')
@login_required
@role_required(['creator', 'director'])
def api_workspace_detail(ws_id):
    workspace = Workspace.query.get_or_404(ws_id)
    if not can_view_workspace(current_user, workspace):
        return _workspace_access_denied_json()
    payload = serialize_workspace_detail(workspace, _admin_avatar_url)
    payload['permissions'] = workspace_panel_permissions(current_user, workspace)
    return jsonify({'status': 'success', 'workspace': payload})


@admin_bp.route('/api/admin/workspaces/<int:ws_id>/extend', methods=['POST'])
@login_required
@role_required(['creator', 'director'])
def api_workspace_extend(ws_id):
    workspace = Workspace.query.get_or_404(ws_id)
    if not can_extend_workspace_subscription(current_user, workspace):
        return jsonify({
            'status': 'error',
            'error': 'Нет прав на изменение подписки',
        }), 403
    data = request.get_json(silent=True) or {}
    raw = data.get('expires_at') if 'expires_at' in data else request.form.get('expires_at')
    if raw is None or (isinstance(raw, str) and not raw.strip()):
        workspace.expires_at = None
    else:
        workspace.expires_at = _parse_expires_at(raw)
    db.session.commit()
    return jsonify({
        'status': 'success',
        'workspace': serialize_workspace_detail(workspace, _admin_avatar_url),
    })


@admin_bp.route('/api/admin/workspaces/<int:ws_id>/sessions/terminate', methods=['POST'])
@login_required
@role_required(['creator', 'director'])
def api_workspace_terminate_session(ws_id):
    workspace = Workspace.query.get_or_404(ws_id)
    if not can_terminate_workspace_admin_sessions(current_user, workspace):
        return jsonify({
            'status': 'error',
            'error': 'Нет прав на управление сеансами',
        }), 403
    data = request.get_json(silent=True) or {}
    user_id = data.get('user_id')
    session_id = data.get('session_id')
    if not user_id or not session_id:
        return jsonify({'status': 'error', 'error': 'Укажите user_id и session_id'}), 400

    user = User.query.filter_by(id=int(user_id), workspace_id=workspace.id).first()
    if not user or user.role not in ADMIN_STAFF_ROLES:
        return jsonify({'status': 'error', 'error': 'Пользователь не найден в админском составе'}), 404

    if not admin_terminate_session(user, int(session_id)):
        return jsonify({'status': 'error', 'error': 'Сеанс не найден'}), 404

    return jsonify({
        'status': 'success',
        'workspace': serialize_workspace_detail(workspace, _admin_avatar_url),
    })


@admin_bp.route('/admin/clients')
@login_required
@role_required(list(CLIENT_ADMIN_ROLES))
def manage_clients():
    ws_id, err = require_workspace(current_user)
    if err:
        flash(err)
        return redirect(url_for('admin.manage_users'))

    clients = (
        Client.query.filter_by(workspace_id=ws_id)
        .order_by(Client.name.asc())
        .all()
    )
    return render_template(
        'admin/clients.html',
        **_admin_context('clients', clients=clients),
    )


@admin_bp.route('/admin/workspace/create', methods=['POST'])
@login_required
@role_required(['creator', 'director'])
def create_workspace():
    if not can_create_workspace(current_user):
        flash('Создание пространств доступно только создателю платформы')
        abort(403)

    name = normalize_workspace_name(request.form.get('name'))
    from utils.workspace_utils import DEFAULT_WORKSPACE_ADMIN_LIMIT, normalize_admin_limit_value

    admin_limit_raw = request.form.get('admin_limit', str(DEFAULT_WORKSPACE_ADMIN_LIMIT))
    expires_at = _parse_expires_at(request.form.get('expires_at'))

    if not name:
        flash('Укажите название рабочего пространства')
        return redirect(url_for('admin.manage_workspaces'))

    admin_limit, limit_err = normalize_admin_limit_value(admin_limit_raw, current_user)
    if limit_err:
        flash(limit_err)
        return redirect(url_for('admin.manage_workspaces'))

    if Workspace.query.filter_by(name=name).first():
        flash('Пространство с таким названием уже существует')
        return redirect(url_for('admin.manage_workspaces'))

    workspace = Workspace(
        name=name,
        admin_limit=admin_limit,
        expires_at=expires_at,
        invite_key=generate_invite_key(),
        created_by_id=current_user.id,
    )
    db.session.add(workspace)
    db.session.commit()
    flash(f'Рабочее пространство «{name}» создано. Ключ директора: {workspace.invite_key}')
    return redirect(url_for('admin.manage_workspaces'))


@admin_bp.route('/admin/workspace/generate-share-link', methods=['POST'])
@login_required
@role_required(['director'])
def generate_share_link():
    ws_id, err = require_workspace(current_user)
    if err:
        flash(err)
        return redirect(url_for('admin.manage_workspaces'))

    workspace = Workspace.query.get_or_404(ws_id)
    workspace.share_token = uuid.uuid4().hex[:16]
    db.session.commit()
    flash('Ссылка для сотрудников сгенерирована')
    return redirect(url_for('admin.manage_workspaces'))


@admin_bp.route('/admin/clients/add', methods=['POST'])
@login_required
@role_required(list(CLIENT_ADMIN_ROLES))
def add_client():
    ws_id, err = require_workspace(current_user)
    if err:
        flash(err)
        return redirect(url_for('admin.manage_clients'))

    name = (request.form.get('name') or '').strip()
    phone = _normalize_phone(request.form.get('phone'))
    address = (request.form.get('address') or '').strip() or None

    if not name:
        flash('Укажите имя или название организации')
        return redirect(url_for('admin.manage_clients'))

    if request.form.get('phone') and not phone:
        flash('Некорректный номер телефона')
        return redirect(url_for('admin.manage_clients'))

    client = Client(
        workspace_id=ws_id,
        name=name,
        phone=encrypt_data(phone) if phone else None,
        address=address,
    )
    db.session.add(client)
    db.session.commit()
    flash(f'Клиент «{name}» добавлен')
    return redirect(url_for('admin.manage_clients'))


@admin_bp.route('/admin/clients/import', methods=['POST'])
@login_required
@role_required(list(CLIENT_ADMIN_ROLES))
def import_clients():
    ws_id, err = require_workspace(current_user)
    if err:
        flash(err)
        return redirect(url_for('admin.manage_clients'))

    file = request.files.get('file')
    if not file or not file.filename:
        flash('Выберите файл Excel (.xlsx)')
        return redirect(url_for('admin.manage_clients'))

    if not file.filename.lower().endswith('.xlsx'):
        flash('Допустим только формат .xlsx')
        return redirect(url_for('admin.manage_clients'))

    try:
        rows, parse_error = _parse_clients_workbook(file)
    except Exception:
        flash('Не удалось прочитать файл Excel')
        return redirect(url_for('admin.manage_clients'))

    if parse_error:
        flash(parse_error)
        return redirect(url_for('admin.manage_clients'))

    if not rows:
        flash('В файле нет строк для импорта')
        return redirect(url_for('admin.manage_clients'))

    added = 0
    for row in rows:
        raw_phone = row.get('phone')
        db.session.add(Client(
            workspace_id=ws_id,
            name=row['name'],
            phone=encrypt_data(raw_phone) if raw_phone else None,
            address=row.get('address'),
        ))
        added += 1
    db.session.commit()
    flash(f'Импортировано клиентов: {added}')
    return redirect(url_for('admin.manage_clients'))


@admin_bp.route('/admin/clients/<int:client_id>/toggle-status', methods=['POST'])
@login_required
@role_required(list(CLIENT_ADMIN_ROLES))
def toggle_client_status(client_id):
    ws_id, err = require_workspace(current_user)
    if err:
        return jsonify({'ok': False, 'error': err}), 403

    client, access_err = _get_workspace_client(client_id, ws_id)
    if access_err:
        return jsonify({'ok': False, 'error': access_err[0]}), access_err[1]

    client.status = (
        CLIENT_STATUS_FROZEN
        if (client.status or CLIENT_STATUS_ACTIVE) != CLIENT_STATUS_FROZEN
        else CLIENT_STATUS_ACTIVE
    )
    db.session.commit()
    return jsonify({'ok': True, 'client': _client_json(client)})


@admin_bp.route('/admin/clients/<int:client_id>/delete', methods=['POST'])
@login_required
@role_required(list(CLIENT_ADMIN_ROLES))
def delete_client(client_id):
    ws_id, err = require_workspace(current_user)
    if err:
        return jsonify({'ok': False, 'error': err}), 403

    client, access_err = _get_workspace_client(client_id, ws_id)
    if access_err:
        return jsonify({'ok': False, 'error': access_err[0]}), access_err[1]

    client_name = client.name
    _delete_client_record(client)
    db.session.commit()
    return jsonify({'ok': True, 'id': client_id, 'message': f'Клиент «{client_name}» удалён'})


@admin_bp.route('/admin/clients/edit/<int:client_id>', methods=['POST'])
@login_required
@role_required(list(CLIENT_ADMIN_ROLES))
def edit_client(client_id):
    ws_id, err = require_workspace(current_user)
    if err:
        flash(err)
        return redirect(url_for('admin.manage_clients'))

    client, access_err = _get_workspace_client(client_id, ws_id)
    if access_err:
        flash(access_err[0])
        return redirect(url_for('admin.manage_clients'))

    name = (request.form.get('name') or '').strip()
    phone = _normalize_phone(request.form.get('phone'))
    address = (request.form.get('address') or '').strip() or None

    if not name:
        flash('Укажите имя или название организации')
        return redirect(url_for('admin.manage_clients'))

    if request.form.get('phone') and not phone:
        flash('Некорректный номер телефона')
        return redirect(url_for('admin.manage_clients'))

    client.name = name
    client.phone = encrypt_data(phone) if phone else None
    client.address = address
    db.session.commit()
    flash(f'Клиент «{name}» обновлён')
    return redirect(url_for('admin.manage_clients'))


@admin_bp.route('/admin/user/<int:user_id>/assign-role', methods=['POST'])
@login_required
@role_required(list(USER_MANAGEMENT_ROLES))
def assign_role(user_id):
    target_user = User.query.get_or_404(user_id)
    ws_id = get_workspace_id(current_user)
    if ws_id and target_user.workspace_id != ws_id:
        flash('Пользователь из другого рабочего пространства')
        return redirect(url_for('admin.manage_users'))

    if target_user.role == 'creator' and current_user.role != 'creator':
        flash('Нельзя изменить роль создателя системы')
        return redirect(url_for('admin.manage_users'))

    if target_user.role == 'director' and current_user.role != 'creator':
        flash('Только создатель системы может менять роль директора')
        return redirect(url_for('admin.manage_users'))

    new_role = request.form.get('role')
    if not current_user.can_assign_role(new_role):
        flash(f'Вы не можете назначать роль {new_role}')
        return redirect(url_for('admin.manage_users'))

    workspace_id = target_user.workspace_id or ws_id
    if workspace_id and would_exceed_admin_limit(
        workspace_id, target_user, new_role, acting_user=current_user
    ):
        flash('Превышен лимит административного состава для этого рабочего пространства')
        return redirect(url_for('admin.manage_users'))

    target_user.role = new_role
    target_user.created_by_id = current_user.id

    rating_raw = request.form.get('rating')
    if rating_raw not in (None, '') and target_user.role in ['worker', 'brigadir']:
        try:
            rating = float(rating_raw)
        except (TypeError, ValueError):
            flash('Рейтинг должен быть числом от 0 до 10')
            return redirect(url_for('admin.manage_users'))
        if rating < 0 or rating > 10:
            flash('Рейтинг должен быть от 0 до 10')
            return redirect(url_for('admin.manage_users'))
        target_user.set_manual_rating(rating)

    db.session.commit()
    flash(f'Роль пользователя {target_user.name} изменена на {new_role}')
    return _admin_redirect_for_role(new_role)


@admin_bp.route('/admin/user/<int:user_id>/set-rating', methods=['POST'])
@login_required
@role_required(['creator', 'director', 'senior_dispatcher', 'manager'])
def set_rating(user_id):
    target_user = User.query.get_or_404(user_id)
    ws_id = get_workspace_id(current_user)
    if ws_id and target_user.workspace_id != ws_id:
        flash('Пользователь из другого рабочего пространства')
        return redirect(url_for('admin.manage_users'))

    if target_user.role not in ['worker', 'brigadir']:
        flash('Рейтинг можно устанавливать только исполнителям и бригадирам')
        return redirect(url_for('admin.manage_users'))

    rating = request.form.get('rating', type=float)
    if rating is None or rating < 0 or rating > 10:
        flash('Рейтинг должен быть от 0 до 10')
        return redirect(url_for('admin.manage_users'))

    target_user.set_manual_rating(rating)
    db.session.commit()
    flash(f'Рейтинг {target_user.name} установлен на {rating}')
    return redirect(url_for('admin.manage_users'))


@admin_bp.route('/admin/user/<int:user_id>/delete', methods=['POST'])
@login_required
@role_required(['creator', 'director'])
def delete_user(user_id):
    target_user = User.query.get_or_404(user_id)
    ws_id = get_workspace_id(current_user)
    if ws_id and target_user.workspace_id != ws_id:
        flash('Пользователь из другого рабочего пространства')
        return redirect(url_for('admin.manage_users'))

    if target_user.role == 'creator':
        flash('Нельзя удалить создателя системы')
        return redirect(url_for('admin.manage_users'))

    db.session.delete(target_user)
    db.session.commit()
    flash(f'Пользователь {target_user.name} удалён')
    return redirect(url_for('admin.manage_users'))


@admin_bp.route('/api/moderation/count-pending')
@login_required
def count_pending_moderations():
    """Счётчик необработанных анкет для бэйджа в меню (по городам диспетчера)."""
    if not _can_view_moderation_badge():
        return jsonify({'error': 'Forbidden'}), 403

    count = _scoped_verification_query(PENDING_VERIFICATION_STATUS).count()
    return jsonify({'count': count})


@admin_bp.route('/admin/verification')
@login_required
def list_verifications():
    if not _can_moderate_verifications():
        flash('Нет доступа')
        if current_user.role == 'director':
            return redirect(url_for('admin.manage_workspaces'))
        return redirect(url_for('admin.manage_users'))

    requests = _scoped_verification_query('pending').all()
    _attach_verification_secure_media(requests)
    return render_template('admin/verifications.html', requests=requests, active_tab='admin')


def _attach_verification_secure_media(requests):
    from utils.secure_media import issue_verification_media_token, verification_file_exists

    for req in requests:
        req.media_passport_url = None
        req.media_selfie_url = None
        if not _can_view_verification_documents(req):
            continue
        if verification_file_exists(req.passport_photo, current_app):
            req.media_passport_url = url_for(
                'admin.verification_secure_media',
                token=issue_verification_media_token(current_user.id, req.id, 'passport'),
            )
        if verification_file_exists(req.selfie_photo, current_app):
            req.media_selfie_url = url_for(
                'admin.verification_secure_media',
                token=issue_verification_media_token(current_user.id, req.id, 'selfie'),
            )


@admin_bp.route('/admin/verification/media/<path:token>')
@login_required
def verification_secure_media(token):
    from utils.secure_media import (
        debug_media_404,
        read_verification_document,
        validate_verification_media_grant,
    )

    token_in_request = (token or '').strip()
    grant = validate_verification_media_grant(token_in_request, current_user.id)
    if not grant:
        from utils.secure_media import MEDIA_GRANT_SESSION_KEY

        session_grants = session.get(MEDIA_GRANT_SESSION_KEY) or {}
        token_in_session = token_in_request in session_grants
        debug_media_404(
            'invalid_or_expired_token',
            token_received=token_in_request[:80] + ('…' if len(token_in_request) > 80 else ''),
            token_length=len(token_in_request),
            token_in_session=token_in_session,
            session_grants_count=len(session_grants),
            user_id=current_user.id,
        )
        abort(404)

    req = VerificationRequest.query.get(grant['req_id'])
    if not req:
        debug_media_404(
            'verification_request_not_found',
            req_id=grant.get('req_id'),
            doc=grant.get('doc'),
        )
        abort(404)

    if not _can_view_verification_documents(req):
        abort(403)

    field = 'passport_photo' if grant['doc'] == 'passport' else 'selfie_photo'
    stored = getattr(req, field, None)
    file_debug = {}
    data, mime = read_verification_document(stored, current_app, debug_out=file_debug)
    if not data:
        debug_media_404(
            file_debug.get('reason', 'decrypt_or_read_failed'),
            req_id=req.id,
            field=field,
            stored_preview=str(stored)[:60] if stored else None,
            **{k: v for k, v in file_debug.items() if k != 'reason'},
        )
        abort(404)

    return Response(
        data,
        mimetype=mime,
        headers={
            'Cache-Control': 'no-store, no-cache, must-revalidate, private',
            'Pragma': 'no-cache',
            'X-Content-Type-Options': 'nosniff',
        },
    )


@admin_bp.route('/admin/verification/clear-media', methods=['POST'])
@login_required
def clear_verification_media():
    from utils.secure_media import clear_verification_media_grants_for_user

    clear_verification_media_grants_for_user(current_user.id)
    return jsonify({'ok': True})


@admin_bp.route('/admin/workers/<int:user_id>/verification-documents')
@login_required
@role_required(list(USER_MANAGEMENT_ROLES))
def worker_verification_documents(user_id):
    from utils.user_profiles import can_view_worker_card
    from utils.secure_media import build_worker_verification_documents_payload, find_verification_documents

    user = User.query.get_or_404(user_id)
    if not can_view_worker_card(current_user, user):
        return jsonify({'ok': False, 'error': 'Нет доступа'}), 403

    ws_id = get_workspace_id(current_user)
    if ws_id and user.workspace_id != ws_id:
        return jsonify({'ok': False, 'error': 'Исполнитель из другого пространства'}), 403

    req, _ = find_verification_documents(user_id, current_app)
    if req and not _can_view_verification_documents(req):
        return jsonify({'ok': False, 'error': 'Нет доступа к документам'}), 403

    payload = build_worker_verification_documents_payload(
        user_id,
        current_user.id,
        current_app,
        lambda token: url_for('admin.verification_secure_media', token=token),
    )
    return jsonify(payload)


@admin_bp.route('/admin/verification/approve/<int:req_id>')
@login_required
def approve_worker(req_id):
    if not _can_moderate_verifications():
        flash('Нет доступа')
        return redirect(url_for('admin.list_verifications'))

    req = VerificationRequest.query.get_or_404(req_id)
    if not _can_moderate_verification(req):
        flash('Заявка из другого рабочего пространства')
        return redirect(url_for('admin.list_verifications'))

    user = User.query.get(req.user_id)
    if user:
        user.is_verified = True
        user.role = 'worker'
        user.name = decrypt_data(req.full_name)
        user.birth_date = req.birth_date
        user.verified_at = datetime.utcnow()
        if req.city_id:
            _append_city_to_allowed_locations(user, req.city_id)
        req.status = 'approved'
        req.moderated_by_id = current_user.id
        req.moderated_at = datetime.utcnow()
        db.session.commit()
        flash(f'Пользователь {user.name} одобрен как исполнитель!')

    return redirect(url_for('admin.list_verifications'))


@admin_bp.route('/admin/verification/reject/<int:req_id>')
@login_required
def reject_worker(req_id):
    if not _can_moderate_verifications():
        flash('Нет доступа')
        return redirect(url_for('admin.list_verifications'))

    req = VerificationRequest.query.get_or_404(req_id)
    if not _can_moderate_verification(req):
        flash('Заявка из другого рабочего пространства')
        return redirect(url_for('admin.list_verifications'))

    req.status = 'rejected'
    req.moderated_by_id = current_user.id
    req.moderated_at = datetime.utcnow()
    db.session.commit()
    flash('Заявка отклонена')
    return redirect(url_for('admin.list_verifications'))


@admin_bp.route('/admin/user/<int:user_id>/locations', methods=['GET'])
@login_required
def get_user_locations(user_id):
    if current_user.role not in ['creator', 'director', 'senior_dispatcher', 'dispatcher']:
        return {'error': 'Нет прав'}, 403

    user = User.query.get_or_404(user_id)
    ws_id = get_workspace_id(current_user)
    if ws_id and user.workspace_id != ws_id:
        return {'error': 'Нет доступа'}, 403

    locations = user.allowed_locations or []
    if isinstance(locations, str):
        import json
        try:
            locations = json.loads(locations) if locations else []
        except Exception:
            locations = []

    from flask import jsonify
    return jsonify({'locations': locations})


@admin_bp.route('/admin/user/<int:user_id>/update-locations', methods=['POST'])
@login_required
def update_user_locations(user_id):
    if current_user.role not in ['creator', 'director', 'senior_dispatcher', 'dispatcher']:
        from flask import jsonify
        return jsonify({'error': 'Нет прав'}), 403

    user = User.query.get_or_404(user_id)
    ws_id = get_workspace_id(current_user)
    if ws_id and user.workspace_id != ws_id:
        from flask import jsonify
        return jsonify({'error': 'Нет доступа'}), 403

    from flask import jsonify
    data = request.get_json()
    locations = data.get('locations', [])
    user.allowed_locations = locations
    db.session.commit()
    return jsonify({'status': 'success', 'message': 'Города обновлены'})
